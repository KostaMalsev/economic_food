"""Extract the household inputs for Limor's empirical ZU method from Excel."""

from argparse import ArgumentParser
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


HERE = Path(__file__).resolve().parent


SOURCE_SHEET = "חישוב סל נורמטיבי "
OUTPUT_COLUMNS = [
    "misparmb", "persons_count", "c3_actual",
    "c30_actual", "c31_actual", "c32_actual", "c33_actual",
    "c34_actual", "c35_actual", "c36_actual", "c37_actual",
    "c38_actual", "c39_actual", "food_actual", "nonfood_actual",
    "FoodNorm-active",
]


def extract(source_path):
    workbook = load_workbook(source_path, read_only=True, data_only=True)
    if SOURCE_SHEET not in workbook.sheetnames:
        raise KeyError(f"Missing source sheet: {SOURCE_SHEET!r}")
    sheet = workbook[SOURCE_SHEET]

    households = {}
    for row in sheet.iter_rows(min_row=2, values_only=True):
        household_id = row[2]
        if household_id is None:
            continue
        values = {
            "misparmb": int(household_id),
            "persons_count": int(row[4]),
            "c3_actual": row[25],
            "c30_actual": row[26],
            "c31_actual": row[27],
            **{f"c{number}_actual": row[26 + number - 30]
               for number in range(32, 40)},
            "food_actual": row[36],
            "nonfood_actual": row[37],
            # Limor's ~500-household selection uses the active cheapest basket.
            "FoodNorm-active": row[22],
        }
        normalized = {key: (0 if value is None and key.startswith("c") else value)
                      for key, value in values.items()}
        if household_id in households:
            if households[household_id] != normalized:
                raise ValueError(f"Conflicting repeated rows for household {household_id}")
        else:
            households[household_id] = normalized

    data = pd.DataFrame(households.values(), columns=OUTPUT_COLUMNS)
    data = data.sort_values("misparmb").reset_index(drop=True)
    if len(data) != 9017:
        raise ValueError(f"Expected 9,017 households, found {len(data):,}")

    component_sum = data[[f"c{i}_actual" for i in range(32, 40)]].sum(axis=1)
    if not component_sum.equals(data["nonfood_actual"]):
        raise ValueError("C32-C39 do not exactly reconcile to nonfood_actual")
    if (data["c30_actual"] + data["c31_actual"] != data["food_actual"]).any():
        raise ValueError("C30+C31 does not exactly reconcile to food_actual")
    return data


def main():
    parser = ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path)
    parser.add_argument("--output", type=Path,
                        default=HERE / "data" / "household_inputs.csv")
    args = parser.parse_args()
    data = extract(args.source)
    data.to_csv(args.output, index=False, float_format="%.10f")
    print(f"Wrote {len(data):,} unique households to {args.output}")


if __name__ == "__main__":
    main()
